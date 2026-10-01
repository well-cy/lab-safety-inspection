# -*- coding: utf-8 -*-
"""
违规记录查询增强 + 导出接口测试（新增独立文件）。

运行：venv/bin/python tests/test_export_api.py
    （Windows: venv\\Scripts\\python tests\\test_export_api.py）

覆盖内容：
  1. GET /api/violations 新增参数与字段
     - total_count（新增）与 total（保持原语义）并存
     - date_from / date_to 日期区间
     - sort / order 排序
  2. GET /api/export/violations 导出
     - fmt=csv  → Content-Type / BOM / 表头 / 行数
     - fmt=xlsx → 可被 openpyxl 正确读取
     - 与 /api/violations 用同一筛选口径（导出条数 == total_count）
  3. 参数校验（非法格式、非法日期、日期区间反向、超限）

测试数据使用 2020-01 的历史日期，避免与真实/演示数据相互干扰；
用例结束在 finally 中清理，不污染演示数据库。

注意：本文件为新增，不修改 tests/run_tests.py 的 T01~T14。
"""
import csv
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from backend.main import app  # noqa: E402

RESULTS = []

# 专用测试日期区间（历史日期，与真实数据隔离）
D1, D2, D3 = "2020-01-01", "2020-01-02", "2020-01-03"


def report(name, ok, detail=""):
    RESULTS.append({"name": name, "ok": ok, "detail": detail})
    print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    if detail:
        print(f"       {detail}")


def seed_test_data():
    """插入 3 条受控违规记录：D1/D2 为「一般违规」，D3 为「严重违规」"""
    from backend import database as db
    conn = db.get_conn()
    ids = []
    for d, sev, types in ((D1, "一般违规", "未佩戴口罩"),
                          (D2, "一般违规", "未佩戴口罩"),
                          (D3, "严重违规", "未佩戴手套,未穿实验服")):
        cur = conn.execute(
            "INSERT INTO detection_records(lab_id, source_type, source_name,"
            " person_count, violation_count, normal_count, process_time_ms,"
            " created_at) VALUES(1,'image','export_api_test',1,1,1,1.0,?)",
            (f"{d} 10:00:00",))
        rid = cur.lastrowid
        conn.execute(
            "INSERT INTO violation_events(record_id, lab_id, area_name, person_id,"
            " violation_types, severity, screenshot_path, frame_time, created_at)"
            " VALUES(?,1,'实验操作区',1,?,?,'',0,?)",
            (rid, types, sev, f"{d} 10:00:00"))
        ids.append(rid)
    conn.commit()
    conn.close()
    return ids


def cleanup(record_ids):
    from backend import database as db
    conn = db.get_conn()
    for rid in record_ids:
        conn.execute("DELETE FROM violation_events WHERE record_id=?", (rid,))
        conn.execute("DELETE FROM detection_records WHERE id=?", (rid,))
    conn.commit()
    conn.close()


def main():
    record_ids = seed_test_data()
    try:
        with TestClient(app) as client:
            print("=" * 64)
            print("违规记录查询增强 + 导出接口测试")
            print("=" * 64)

            # ---------- 1. 查询增强 ----------
            r = client.get(f"/api/violations?date_from={D1}&date_to={D3}&limit=200")
            body = r.json()
            report("查询 - 日期区间过滤",
                   r.status_code == 200 and body.get("total_count") == 3,
                   f"HTTP {r.status_code} total_count={body.get('total_count')} (期望 3)")
            report("查询 - total 保持原语义（本次返回条数）",
                   body.get("total") == len(body.get("items", [])),
                   f"total={body.get('total')} items={len(body.get('items', []))}")

            r = client.get(f"/api/violations?date_from={D2}&date_to={D3}")
            report("查询 - 区间起点过滤（D2 起应为 2 条）",
                   r.json().get("total_count") == 2,
                   f"total_count={r.json().get('total_count')}")

            r = client.get(f"/api/violations?date_from={D1}&date_to={D3}&severity=严重违规")
            report("查询 - severity 精确过滤（应为 1 条）",
                   r.json().get("total_count") == 1,
                   f"total_count={r.json().get('total_count')}")

            r = client.get(f"/api/violations?date_from={D1}&date_to={D3}&vtype=未佩戴口罩")
            report("查询 - vtype 模糊过滤（应为 2 条）",
                   r.json().get("total_count") == 2,
                   f"total_count={r.json().get('total_count')}")

            r = client.get(f"/api/violations?date_from={D1}&date_to={D3}"
                           f"&sort=created_at&order=asc")
            items = r.json().get("items", [])
            ok = len(items) >= 2 and items[0]["created_at"] < items[-1]["created_at"]
            report("查询 - sort=created_at&order=asc 生效", ok,
                   f"首条={items[0]['created_at'] if items else '-'} "
                   f"末条={items[-1]['created_at'] if items else '-'}")

            # ---------- 2. CSV 导出 ----------
            r = client.get(f"/api/export/violations?fmt=csv"
                           f"&date_from={D1}&date_to={D3}")
            ok = (r.status_code == 200
                  and "text/csv" in r.headers.get("content-type", ""))
            report("导出 - CSV 状态码与 Content-Type", ok,
                   f"HTTP {r.status_code} type={r.headers.get('content-type')}")

            raw = r.content
            report("导出 - CSV 带 UTF-8 BOM（Excel 中文不乱码）",
                   raw.startswith(b"\xef\xbb\xbf"),
                   f"前 3 字节={raw[:3]!r}")

            text = raw.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text)))
            header = rows[0]
            report("导出 - CSV 表头正确",
                   header[:4] == ["记录ID", "违规时间", "实验室", "区域"],
                   f"表头={header}")
            report("导出 - CSV 行数与筛选一致（3 条）", len(rows) - 1 == 3,
                   f"数据行={len(rows) - 1}")

            r = client.get(f"/api/export/violations?fmt=csv"
                           f"&date_from={D1}&date_to={D3}&severity=严重违规")
            rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
            ok = len(rows) - 1 == 1 and rows[1][6] == "严重违规"
            report("导出 - CSV 按 severity 过滤生效", ok,
                   f"数据行={len(rows) - 1} severity={rows[1][6] if len(rows) > 1 else '-'}")

            # ---------- 3. XLSX 导出 ----------
            r = client.get(f"/api/export/violations?fmt=xlsx"
                           f"&date_from={D1}&date_to={D3}")
            ok = (r.status_code == 200
                  and "spreadsheetml" in r.headers.get("content-type", ""))
            report("导出 - XLSX 状态码与 Content-Type", ok,
                   f"HTTP {r.status_code} type={r.headers.get('content-type')}")

            try:
                wb = load_workbook(io.BytesIO(r.content))
                ws = wb.active
                heads = [c.value for c in ws[1]]
                report("导出 - XLSX 可被 openpyxl 读取",
                       heads[:4] == ["记录ID", "违规时间", "实验室", "区域"],
                       f"工作表={ws.title} 表头={heads[:4]}")
                report("导出 - XLSX 行数与筛选一致（3 条）",
                       ws.max_row - 1 == 3, f"数据行={ws.max_row - 1}")
            except Exception as exc:  # pragma: no cover
                report("导出 - XLSX 可被 openpyxl 读取", False, f"异常: {exc}")

            # ---------- 4. 导出与页面查询口径一致性 ----------
            q = client.get(f"/api/violations?date_from={D1}&date_to={D3}&limit=1").json()
            e = client.get(f"/api/export/violations?fmt=csv"
                           f"&date_from={D1}&date_to={D3}")
            exported = len(list(csv.reader(
                io.StringIO(e.content.decode("utf-8-sig"))))) - 1
            report("一致性 - 导出条数 == 查询 total_count",
                   exported == q["total_count"],
                   f"导出={exported} total_count={q['total_count']}")

            report("一致性 - X-Total-Count 响应头正确",
                   e.headers.get("x-total-count") == str(q["total_count"]),
                   f"X-Total-Count={e.headers.get('x-total-count')}")

            # ---------- 5. 参数校验 ----------
            cases = [
                ("导出 - 非法格式 fmt=pdf（422）",
                 "/api/export/violations?fmt=pdf", 422),
                ("导出 - 非法日期 2026-13-99（400）",
                 "/api/export/violations?fmt=csv&date_from=2026-13-99", 400),
                ("导出 - 不存在的日期 2026-02-30（400）",
                 "/api/export/violations?fmt=csv&date_from=2026-02-30", 400),
                ("导出 - 非日期字符串 abc（400）",
                 "/api/export/violations?fmt=csv&date_from=abc", 400),
                ("导出 - 日期区间反向（400）",
                 "/api/export/violations?fmt=csv&date_from=2026-10-05&date_to=2026-10-01", 400),
                ("导出 - 非法 order（422）",
                 "/api/export/violations?fmt=csv&order=random", 422),
                ("导出 - lab_id=0 越界（422）",
                 "/api/export/violations?fmt=csv&lab_id=0", 422),
            ]
            for name, url, expect in cases:
                rr = client.get(url)
                report(name, rr.status_code == expect,
                       f"HTTP {rr.status_code}（期望 {expect}）")

    finally:
        cleanup(record_ids)

    passed = sum(1 for x in RESULTS if x["ok"])
    total = len(RESULTS)
    print("=" * 64)
    print(f"结果: {passed}/{total} 通过")
    if passed != total:
        print("失败用例:")
        for x in RESULTS:
            if not x["ok"]:
                print(f"  - {x['name']}: {x['detail']}")
    print("=" * 64)
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
