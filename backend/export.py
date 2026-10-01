# -*- coding: utf-8 -*-
"""
违规记录导出模块（独立文件）。

前端契约来源：backend/static/export.js（成员 5 已写好的占位文件）
    GET /api/export/violations?fmt=csv|xlsx&date_from=&date_to=&severity=&ppe_type=

设计约定：
- 只**新增**端点，不改动既有 GET /api/violations 的路径与响应字段
- 筛选口径复用 backend/database.py 的 _violation_where，保证"页面看到的"与
  "导出的"完全一致（同一次筛选不会导出不同结果）
- CSV 用 utf-8-sig（带 BOM）编码，Excel 双击打开中文不乱码
- 导出文件名保持纯 ASCII（规避中文路径在部分环境下的问题）

为什么要独立成文件：
    docs/PROJECT_FRAMEWORK_GUIDE.md 第 5 节要求"独立功能必须有独立文件"，
    避免与其他成员的改动在 backend/main.py 上产生冲突。
"""
import csv
import io
from datetime import datetime

from fastapi import APIRouter, Query, Response
from fastapi.responses import JSONResponse

from backend import database as db
from backend.validators import is_valid_date

router = APIRouter(prefix="/api/export", tags=["违规记录导出"])

# 导出列：(数据库字段, 中文表头)
EXPORT_COLUMNS = [
    ("id", "记录ID"),
    ("created_at", "违规时间"),
    ("lab_name", "实验室"),
    ("area_name", "区域"),
    ("person_id", "人员编号"),
    ("violation_types", "违规类型"),
    ("severity", "严重程度"),
    ("frame_time", "视频内时间(秒)"),
    ("screenshot_path", "截图文件"),
]

# 单次导出上限，防止误操作把整库拉出来占满内存
MAX_EXPORT_ROWS = 50000

XLSX_MEDIA_TYPE = ("application/vnd.openxmlformats-officedocument"
                   ".spreadsheetml.sheet")

def _build_csv(rows: list) -> bytes:
    """生成 CSV 字节流（utf-8-sig 带 BOM，Excel 打开中文不乱码）"""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow([cn for _, cn in EXPORT_COLUMNS])
    for row in rows:
        writer.writerow(["" if row.get(k) is None else row.get(k)
                         for k, _ in EXPORT_COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


def _build_xlsx(rows: list) -> bytes:
    """生成 xlsx 字节流（使用 openpyxl）"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "违规记录"

    ws.append([cn for _, cn in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(["" if row.get(k) is None else row.get(k)
                   for k, _ in EXPORT_COLUMNS])

    # 按中文表头长度粗略设置列宽
    for idx, (_, cn) in enumerate(EXPORT_COLUMNS, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = max(len(cn) * 2.2, 12)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/violations")
def export_violations(
    fmt: str = Query("csv", pattern="^(csv|xlsx)$",
                     description="导出格式：csv 或 xlsx"),
    date_from: str | None = Query(None, description="起始日期 YYYY-MM-DD（含）"),
    date_to: str | None = Query(None, description="结束日期 YYYY-MM-DD（含）"),
    severity: str | None = Query(None, description="严重程度精确匹配"),
    ppe_type: str | None = Query(None, description="违规类型模糊匹配，如 未佩戴口罩"),
    vtype: str | None = Query(None, description="ppe_type 的别名（兼容保留）"),
    lab_id: int | None = Query(None, ge=1, description="按实验室过滤"),
    sort: str = Query("created_at", description="id|created_at|severity|area_name"),
    order: str = Query("desc", pattern="^(asc|desc)$", description="asc 或 desc"),
):
    """
    导出违规记录（CSV / Excel）。

    与 GET /api/violations 使用同一套筛选口径；不传 severity 时同样默认
    排除"正常"评估记录，保证导出内容与页面所见表一致。
    """
    keyword = ppe_type or vtype

    # 日期校验：格式与真实性都检查，统一返回 400 + error 字段
    for label, value in (("date_from", date_from), ("date_to", date_to)):
        if not is_valid_date(value):
            return JSONResponse(
                {"error": f"{label} 不是有效日期：{value}（应为 YYYY-MM-DD）"},
                status_code=400)

    if date_from and date_to and date_from > date_to:
        return JSONResponse(
            {"error": f"日期区间无效：起始 {date_from} 晚于结束 {date_to}"},
            status_code=400)

    total = db.count_violations(vtype=keyword, severity=severity,
                                date_from=date_from, date_to=date_to,
                                lab_id=lab_id)
    if total > MAX_EXPORT_ROWS:
        return JSONResponse(
            {"error": f"符合条件记录 {total} 条，超过单次导出上限 "
                      f"{MAX_EXPORT_ROWS} 条，请缩小日期或筛选范围"},
            status_code=400)

    rows = db.query_violations(vtype=keyword, severity=severity,
                               date_from=date_from, date_to=date_to,
                               lab_id=lab_id, limit=max(total, 1), offset=0,
                               sort=sort, order=order)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if fmt == "xlsx":
        content = _build_xlsx(rows)
        media_type = XLSX_MEDIA_TYPE
        filename = f"violations_{stamp}.xlsx"
    else:
        content = _build_csv(rows)
        media_type = "text/csv; charset=utf-8"
        filename = f"violations_{stamp}.csv"

    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Total-Count": str(total),   # 便于前端显示"共导出 N 条"
        },
    )
